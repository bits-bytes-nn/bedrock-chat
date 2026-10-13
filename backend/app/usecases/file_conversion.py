"""Convert files that Bedrock cannot read directly into text.

Presentations are converted synchronously. Audio and video are transcribed with
Amazon Transcribe, which runs asynchronously: callers start the conversion and poll
it until it completes. The text is written next to the source object as
`<filename>.txt`.
"""

import hashlib
import io
import json
import logging
import os
from pathlib import PurePosixPath
from typing import Literal

import boto3
from botocore.exceptions import ClientError
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pydantic import BaseModel
from ulid import ULID

logger = logging.getLogger(__name__)

DOCUMENT_BUCKET = os.environ.get("DOCUMENT_BUCKET", "bedrock-documents")
BEDROCK_REGION = os.environ.get("BEDROCK_REGION", "us-east-1")

PRESENTATION_EXTENSIONS = frozenset({".pptx"})
# Media formats supported by Amazon Transcribe batch jobs.
MEDIA_EXTENSIONS = frozenset(
    {".mp3", ".mp4", ".wav", ".flac", ".ogg", ".amr", ".webm", ".m4a"}
)
CONVERTIBLE_EXTENSIONS = PRESENTATION_EXTENSIONS | MEDIA_EXTENSIONS
# Presentations are parsed in the API Lambda, so their size is limited. Recordings
# are limited because transcription is billed per minute.
MAX_PRESENTATION_BYTES = 100 * 1024 * 1024
MAX_MEDIA_BYTES = 500 * 1024 * 1024
# Languages that Transcribe chooses from when identifying the spoken language.
TRANSCRIPTION_LANGUAGES = ["ko-KR", "en-US", "ja-JP", "zh-CN"]
MAX_SPEAKERS = 10
# Prefix of the jobs this application starts; the IAM policy is scoped to it.
TRANSCRIPTION_JOB_PREFIX = "brchat-"

s3_client = boto3.client("s3", BEDROCK_REGION)
transcribe_client = boto3.client("transcribe", BEDROCK_REGION)


class ConversionStatus(BaseModel):
    status: Literal["IN_PROGRESS", "COMPLETED", "FAILED"]
    reason: str = ""


def is_convertible(filename: str) -> bool:
    return PurePosixPath(filename).suffix.lower() in CONVERTIBLE_EXTENSIONS


def chat_upload_key(user_id: str, upload_id: str, filename: str) -> str:
    """Key of a chat attachment to convert. `.temp/` objects expire after a day.

    Raises:
        ValueError: If the upload id or the filename is invalid.
    """
    try:
        ULID.from_str(upload_id)
    except ValueError:
        raise ValueError("Upload ID must be a ULID.") from None
    if not filename or "/" in filename or not is_convertible(filename):
        raise ValueError("Invalid filename.")
    return f".temp/{user_id}/{upload_id}/{filename}"


def converted_key(source_key: str) -> str:
    return f"{source_key}.txt"


def _source(source_key: str) -> dict:
    """Return the metadata of the uploaded object.

    Raises:
        ValueError: If nothing was uploaded at the key.
    """
    try:
        return s3_client.head_object(Bucket=DOCUMENT_BUCKET, Key=source_key)
    except ClientError as e:
        if e.response["Error"]["Code"] in ("404", "NoSuchKey"):
            raise ValueError("The file was not uploaded.") from None
        raise


def _delete_if_older(key: str, than) -> None:
    try:
        stale = s3_client.head_object(Bucket=DOCUMENT_BUCKET, Key=key)
    except ClientError as e:
        if e.response["Error"]["Code"] in ("404", "NoSuchKey"):
            return
        raise
    # S3 times have one-second precision. Results written in the same second as
    # the upload are older than it: presentations are converted again anyway.
    if stale["LastModified"] <= than:
        s3_client.delete_object(Bucket=DOCUMENT_BUCKET, Key=key)


def _job_name(source_key: str, source: dict) -> str:
    # Deterministic per upload, so that starting a conversion twice reuses the job,
    # while uploading again, even the same bytes, starts a new one.
    upload = f"{source_key}:{source['ETag']}:{source['LastModified'].isoformat()}"
    return TRANSCRIPTION_JOB_PREFIX + hashlib.sha256(upload.encode()).hexdigest()


def _transcript_key(source_key: str, job_name: str) -> str:
    # Transcribe accepts only ASCII output keys, so the job name is used instead of
    # the filename, which may be Korean.
    return str(PurePosixPath(source_key).parent / f"{job_name}.json")


def _exists(key: str) -> bool:
    try:
        s3_client.head_object(Bucket=DOCUMENT_BUCKET, Key=key)
        return True
    except ClientError as e:
        if e.response["Error"]["Code"] in ("404", "NoSuchKey"):
            return False
        raise


def _shape_lines(shapes) -> list[str]:
    lines = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_lines(shape.shapes))
        if shape.has_text_frame and shape.text_frame.text.strip():
            lines.append(shape.text_frame.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return lines


def presentation_to_text(data: bytes) -> str:
    """Extract the text of every slide, including tables and speaker notes."""
    presentation = Presentation(io.BytesIO(data))
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines = [f"## Slide {number}", *_shape_lines(slide.shapes)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        slides.append("\n".join(lines))
    return "\n\n".join(slides)


def transcript_to_text(transcript: dict) -> str:
    """Format a Transcribe result, one line per speaker turn when labeled."""
    results = transcript["results"]
    segments = results.get("audio_segments", [])
    if not segments or not any("speaker_label" in s for s in segments):
        return results["transcripts"][0]["transcript"]

    lines: list[str] = []
    speaker = None
    for segment in segments:
        label = segment.get("speaker_label", speaker)
        text = segment["transcript"].strip()
        if label == speaker and lines:
            lines[-1] = f"{lines[-1]} {text}"
        else:
            # spk_0 -> Speaker 1
            number = int(label.split("_")[-1]) + 1 if label else 1
            lines.append(f"Speaker {number}: {text}")
            speaker = label
    return "\n".join(lines)


def start_conversion(source_key: str) -> ConversionStatus:
    """Start converting the object at `source_key` into text.

    Raises:
        ValueError: If the file type cannot be converted.
    """
    suffix = PurePosixPath(source_key).suffix.lower()
    if suffix not in CONVERTIBLE_EXTENSIONS:
        raise ValueError(f"Files of type {suffix} cannot be converted.")
    source = _source(source_key)
    # A previous upload with the same name may have left its text behind. The text of
    # this upload is kept, so that starting twice is harmless.
    _delete_if_older(converted_key(source_key), source["LastModified"])

    if suffix in PRESENTATION_EXTENSIONS:
        if source["ContentLength"] > MAX_PRESENTATION_BYTES:
            raise ValueError("Presentations must be 100 MB or smaller.")
        data = s3_client.get_object(Bucket=DOCUMENT_BUCKET, Key=source_key)[
            "Body"
        ].read()
        try:
            text = presentation_to_text(data)
        except Exception as e:  # python-pptx raises several types for broken files
            raise ValueError("The presentation could not be read.") from e
        s3_client.put_object(
            Bucket=DOCUMENT_BUCKET,
            Key=converted_key(source_key),
            Body=text.encode(),
            ContentType="text/plain; charset=utf-8",
        )
        return ConversionStatus(status="COMPLETED")
    if source["ContentLength"] > MAX_MEDIA_BYTES:
        raise ValueError("Recordings must be 500 MB or smaller.")
    job_name = _job_name(source_key, source)
    try:
        transcribe_client.start_transcription_job(
            TranscriptionJobName=job_name,
            Media={"MediaFileUri": f"s3://{DOCUMENT_BUCKET}/{source_key}"},
            OutputBucketName=DOCUMENT_BUCKET,
            OutputKey=_transcript_key(source_key, job_name),
            IdentifyLanguage=True,
            LanguageOptions=TRANSCRIPTION_LANGUAGES,
            Settings={"ShowSpeakerLabels": True, "MaxSpeakerLabels": MAX_SPEAKERS},
        )
    except transcribe_client.exceptions.ConflictException:
        logger.info("Transcription job for %s already exists", source_key)
    return ConversionStatus(status="IN_PROGRESS")


def get_conversion(source_key: str) -> ConversionStatus:
    """Return the status of a conversion, finishing it when the transcript is ready."""
    if _exists(converted_key(source_key)):
        return ConversionStatus(status="COMPLETED")
    if PurePosixPath(source_key).suffix.lower() not in MEDIA_EXTENSIONS:
        return ConversionStatus(
            status="FAILED", reason="The conversion was not started."
        )

    job_name = _job_name(source_key, _source(source_key))
    try:
        job = transcribe_client.get_transcription_job(TranscriptionJobName=job_name)[
            "TranscriptionJob"
        ]
    except transcribe_client.exceptions.BadRequestException:
        return ConversionStatus(
            status="FAILED", reason="The conversion was not started."
        )

    status = job["TranscriptionJobStatus"]
    if status == "FAILED":
        return ConversionStatus(status="FAILED", reason=job.get("FailureReason", ""))
    if status != "COMPLETED":
        return ConversionStatus(status="IN_PROGRESS")

    try:
        transcript = json.loads(
            s3_client.get_object(
                Bucket=DOCUMENT_BUCKET, Key=_transcript_key(source_key, job_name)
            )["Body"].read()
        )
    except s3_client.exceptions.NoSuchKey:
        # Saving a bot purges its temporary files, including the transcript.
        return ConversionStatus(
            status="FAILED", reason="The transcript expired. Upload the file again."
        )
    text = transcript_to_text(transcript)
    if not text.strip():
        return ConversionStatus(status="FAILED", reason="No speech was detected.")
    s3_client.put_object(
        Bucket=DOCUMENT_BUCKET,
        Key=converted_key(source_key),
        Body=text.encode(),
        ContentType="text/plain; charset=utf-8",
    )
    return ConversionStatus(status="COMPLETED")


def read_converted_text(source_key: str) -> str:
    return (
        s3_client.get_object(Bucket=DOCUMENT_BUCKET, Key=converted_key(source_key))[
            "Body"
        ]
        .read()
        .decode()
    )

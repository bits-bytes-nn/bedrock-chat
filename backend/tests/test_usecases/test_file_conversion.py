import io
import json
import unittest
from unittest.mock import MagicMock, patch

import boto3
import pytest
from app.usecases import file_conversion
from moto import mock_aws
from pptx import Presentation
from pptx.util import Inches

BUCKET = "test-document-bucket"


def presentation() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "분기 실적"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    table.table.cell(0, 0).text = "Q1"
    table.table.cell(0, 1).text = "120"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1)).text = (
        "grouped note"
    )
    slide.notes_slide.notes_text_frame.text = "발표자 노트"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


class TestPresentationToText(unittest.TestCase):
    def test_extracts_titles_tables_groups_and_notes(self):
        text = file_conversion.presentation_to_text(presentation())

        self.assertIn("## Slide 1", text)
        self.assertIn("분기 실적", text)
        self.assertIn("Q1 | 120", text)
        self.assertIn("grouped note", text)
        self.assertIn("Notes: 발표자 노트", text)


class TestTranscriptToText(unittest.TestCase):
    def test_merges_turns_by_speaker(self):
        transcript = {
            "results": {
                "transcripts": [{"transcript": "all"}],
                "audio_segments": [
                    {"speaker_label": "spk_0", "transcript": "안녕하세요."},
                    {"speaker_label": "spk_0", "transcript": "회의를 시작합니다."},
                    {"speaker_label": "spk_1", "transcript": "네."},
                ],
            }
        }
        self.assertEqual(
            file_conversion.transcript_to_text(transcript),
            "Speaker 1: 안녕하세요. 회의를 시작합니다.\nSpeaker 2: 네.",
        )

    def test_uses_the_plain_transcript_without_speakers(self):
        transcript = {"results": {"transcripts": [{"transcript": "hello"}]}}
        self.assertEqual(file_conversion.transcript_to_text(transcript), "hello")


class TestChatUploadKey(unittest.TestCase):
    def test_builds_an_expiring_key(self):
        key = file_conversion.chat_upload_key(
            "user", "01J9Z3Q4R5S6T7V8W9X0Y1Z2A3", "slides.pptx"
        )
        self.assertEqual(key, ".temp/user/01J9Z3Q4R5S6T7V8W9X0Y1Z2A3/slides.pptx")

    def test_rejects_invalid_input(self):
        for upload_id, filename in (
            ("../x", "slides.pptx"),
            ("01J9Z3Q4R5S6T7V8W9X0Y1Z2A3", "a/b.pptx"),
            ("01J9Z3Q4R5S6T7V8W9X0Y1Z2A3", "notes.txt"),
        ):
            with self.assertRaises(ValueError):
                file_conversion.chat_upload_key("user", upload_id, filename)


@pytest.fixture
def document_bucket():
    with mock_aws():
        s3 = boto3.client("s3", region_name="us-east-1")
        s3.create_bucket(Bucket=BUCKET)
        with patch.object(file_conversion, "DOCUMENT_BUCKET", BUCKET), patch.object(
            file_conversion, "s3_client", s3
        ):
            yield s3


@pytest.mark.usefixtures("document_bucket")
class TestConversion(unittest.TestCase):
    def setUp(self):
        self.s3 = file_conversion.s3_client

    def test_converts_presentations_immediately(self):
        self.s3.put_object(
            Bucket=BUCKET, Key="u/b/_temp/deck.pptx", Body=presentation()
        )

        status = file_conversion.start_conversion("u/b/_temp/deck.pptx")

        self.assertEqual(status.status, "COMPLETED")
        self.assertIn(
            "분기 실적", file_conversion.read_converted_text("u/b/_temp/deck.pptx")
        )

    def test_rejects_broken_presentations(self):
        self.s3.put_object(Bucket=BUCKET, Key="u/b/_temp/deck.pptx", Body=b"not a zip")
        with self.assertRaises(ValueError):
            file_conversion.start_conversion("u/b/_temp/deck.pptx")

    def test_transcribes_recordings(self):
        key = "u/b/_temp/meeting.m4a"
        self.s3.put_object(Bucket=BUCKET, Key=key, Body=b"audio")
        transcribe = MagicMock()
        transcribe.exceptions.ConflictException = type("Conflict", (Exception,), {})
        transcribe.exceptions.BadRequestException = type("BadRequest", (Exception,), {})
        with patch.object(file_conversion, "transcribe_client", transcribe):
            self.assertEqual(
                file_conversion.start_conversion(key).status, "IN_PROGRESS"
            )
            job = transcribe.start_transcription_job.call_args.kwargs
            self.assertEqual(job["Media"]["MediaFileUri"], f"s3://{BUCKET}/{key}")
            self.assertTrue(job["TranscriptionJobName"].startswith("brchat-"))

            transcribe.get_transcription_job.return_value = {
                "TranscriptionJob": {"TranscriptionJobStatus": "IN_PROGRESS"}
            }
            self.assertEqual(file_conversion.get_conversion(key).status, "IN_PROGRESS")

            self.s3.put_object(
                Bucket=BUCKET,
                Key=job["OutputKey"],
                Body=json.dumps({"results": {"transcripts": [{"transcript": "회의"}]}}),
            )
            transcribe.get_transcription_job.return_value = {
                "TranscriptionJob": {"TranscriptionJobStatus": "COMPLETED"}
            }
            self.assertEqual(file_conversion.get_conversion(key).status, "COMPLETED")

        self.assertEqual(file_conversion.read_converted_text(key), "회의")

    def test_new_uploads_with_the_same_name_start_a_new_job(self):
        key = "u/b/_temp/meeting.mp3"
        transcribe = MagicMock()
        transcribe.exceptions.ConflictException = type("Conflict", (Exception,), {})
        with patch.object(file_conversion, "transcribe_client", transcribe):
            self.s3.put_object(Bucket=BUCKET, Key=key, Body=b"first")
            file_conversion.start_conversion(key)
            self.s3.put_object(Bucket=BUCKET, Key=f"{key}.txt", Body=b"old text")
            self.s3.put_object(Bucket=BUCKET, Key=key, Body=b"second")
            file_conversion.start_conversion(key)
            transcribe.get_transcription_job.return_value = {
                "TranscriptionJob": {"TranscriptionJobStatus": "IN_PROGRESS"}
            }
            # The text of the previous upload is not reported for the new one.
            self.assertEqual(file_conversion.get_conversion(key).status, "IN_PROGRESS")

        first, second = [
            c.kwargs["TranscriptionJobName"]
            for c in transcribe.start_transcription_job.call_args_list
        ]
        self.assertNotEqual(first, second)

    def test_rejects_large_presentations(self):
        key = "u/b/_temp/deck.pptx"
        self.s3.put_object(Bucket=BUCKET, Key=key, Body=presentation())
        with patch.object(file_conversion, "MAX_PRESENTATION_BYTES", 10):
            with self.assertRaises(ValueError):
                file_conversion.start_conversion(key)

    def test_rejects_missing_uploads(self):
        with self.assertRaises(ValueError):
            file_conversion.start_conversion("u/b/_temp/missing.pptx")

    def test_reports_failed_transcriptions(self):
        self.s3.put_object(Bucket=BUCKET, Key="u/b/_temp/meeting.mp3", Body=b"audio")
        transcribe = MagicMock()
        transcribe.exceptions.BadRequestException = type("BadRequest", (Exception,), {})
        transcribe.get_transcription_job.return_value = {
            "TranscriptionJob": {
                "TranscriptionJobStatus": "FAILED",
                "FailureReason": "Unsupported media",
            }
        }
        with patch.object(file_conversion, "transcribe_client", transcribe):
            status = file_conversion.get_conversion("u/b/_temp/meeting.mp3")
        self.assertEqual(
            (status.status, status.reason), ("FAILED", "Unsupported media")
        )


if __name__ == "__main__":
    unittest.main()
